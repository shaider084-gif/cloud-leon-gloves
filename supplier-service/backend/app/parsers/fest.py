"""Поставщик ФЭСТ (ООО «Предприятие «ФЭСТ», аптечки; slug fest) — прайс .xlsx.

Один лист «TDSheet»: строка-раздел («01. Аптечки автомобильные приказ 260н»)
в первой колонке, затем товары: B — артикул, C — номенклатура (первая строка
— название, далее «штрихкод …»), J — ставка НДС, K — цена, M — количество в
коробке, N — «объём м3 / вес кг» коробки. Шапка таблицы повторяется на
каждой странице — такие строки пропускаются.

Файл из учётной системы нестандартный: общие строки лежат в
`xl/SharedStrings.xml` (с заглавной S), и openpyxl его не находит —
переименовываем запись в памяти перед чтением.

Сопоставление с «Весь каталог» по артикулу для этого поставщика НЕ
делается (решение пользователя, 2026-10-05): берутся только данные прайса.
Единственное добавление — «ID варианта» из таблицы пользователя
(fest_ids.json: артикул прайса -> ID варианта в магазине).
Цена из прайса (по артикулу) записывается как закупочная; цена продажи =
закупка × 1,2 (MARKUP) без копеек (отбрасываются, не округляются). Встроенные в файл картинки (≈235) не
используются; фото, описания, параметры и новые названия — из fest_extra.json
(донор promza.ru, см. _load_extra).
"""
import io
import json
import os
import re
import zipfile
from typing import Dict, List, Optional

import openpyxl

from ..schemas import ProductIn
from .base import BaseParser

BRAND = "ФЭСТ"
MARKUP = 1.2  # цена продажи = закупка × 1,2 (решение пользователя, 2026-10-05)


def _load_variant_ids() -> Dict[str, str]:
    """Артикул прайса -> ID варианта в магазине (первая страница файла
    ФЭСТ.xlsx пользователя; 202 из 204 ID есть в «Весь каталог» как «ID варианта»)."""
    path = os.path.join(os.path.dirname(__file__), "fest_ids.json")
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f).get("variant_id_by_article", {})
    except (OSError, ValueError):
        return {}


def _load_extra() -> Dict[str, dict]:
    """Артикул прайса -> {name, images, description, attributes} из донора promza.ru
    (fest_extra.json; собрано из таблицы донора и страниц сайта, описания и параметры
    написаны заново; названия — по шаблону «Товар Бренд описание, приказ N, арт. X»).
    Донор нашёл 178 из 234 артикулов; у остальных — только новое название."""
    path = os.path.join(os.path.dirname(__file__), "fest_extra.json")
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


_VARIANT_IDS = _load_variant_ids()
_EXTRA = _load_extra()

# Реестровые номера (проверено по реестрам 2026-10-06; ООО «Предприятие «ФЭСТ», ИНН 4442016903).
# Росаккредитация (pub.fsa.gov.ru), только ДЕЙСТВУЮЩИЕ документы, где продукт назван прямо:
#  - сертификат на жилет сигнальный (модель 2070, «ФЭСТ») — наборы автомобилиста с жилетом;
#  - декларация на сумки/футляры торговой марки «ФЭСТ» — санитарные сумки и сумка-трансформер.
# «Реестр Минпромторг» (ГИСП, ПП РФ 719): записей ФЭСТ не найдено — оставляем пустым.
_CERT_VEST = "ЕАЭС RU С-RU.ПФ02.В.09123/24"
_DECL_BAGS = "ЕАЭС N RU Д-RU.РА08.В.76884/24"
_REGISTRY = {
    **{a: {"Реестр сертификатов": _CERT_VEST} for a in ("3877", "1463", "1460", "1472", "3879", "1492", "3878", "3875")},
    **{a: {"Реестр сертификатов": _DECL_BAGS} for a in ("1548", "1553", "3160")},
}
COL_SECTION, COL_ARTICLE, COL_NAME, COL_VAT, COL_PRICE, COL_PACK, COL_BOX = 0, 1, 2, 9, 10, 12, 13


def _open_workbook(file_path: str):
    try:
        return openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except KeyError:  # нет xl/sharedStrings.xml — в файле он называется SharedStrings.xml
        pass
    buf = io.BytesIO()
    with zipfile.ZipFile(file_path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data, name = zin.read(item.filename), item.filename
            if name.lower() == "xl/sharedstrings.xml":
                name = "xl/sharedStrings.xml"
            elif name == "[Content_Types].xml":
                data = re.sub(rb"/xl/sharedstrings\.xml", b"/xl/sharedStrings.xml", data, flags=re.I)
            zout.writestr(name, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf, read_only=True, data_only=True)


_TYPE_BY_SECTION = [  # раздел прайса (начало) -> «Тип аптечки»
    ("Аптечки автомобильные", "Автомобильная"),
    ("Аптечки дорожные", "Автомобильная"),
    ("Аптечки транспортные", "Автомобильная"),
    ("Аптечка для оказания работникам", "Производственная"),
    ("Приказы от", "Ведомственная"),
    ("Аптечки для учреждений и производств", "Производственная"),
    ("Аптечки для быта", "Бытовая"),
    ("Медицинское имущество для ГО и ЧС", "ГО и ЧС"),
    ("КИМГЗ", "ГО и ЧС"),
    ("Аптечки отраслевые", "Отраслевая"),
    ("Аптечки по приказам", "Ведомственная"),
    ("Аптечки тактические", "Тактическая"),
]


def _rzn_number(name: str, tu: str = "") -> Optional[str]:
    """Номер регистрационного удостоверения (Росздравнадзор, реестр медицинских изделий,
    elk.roszdravnadzor.gov.ru; проверено 2026-10-06; только действующие РУ ООО «Предприятие «ФЭСТ»).
    Привязка по названию изделия (и ТУ для КИМГЗ), если РУ названо прямо; иначе — пусто."""
    n = (name or "").lower()
    if "кимгз" in n or "гражданской защиты" in n:
        if "161" in (tu or "") or re.search(r"защиты\s*№\s*(1|11|2|3|4|5|7|8|9)\b", n):
            return "РЗН 2018/6809"
        return "РЗН 2020/11903"
    if n.startswith("набор автомобилиста"):
        return None
    rules = [
        ("санитарной сумки", "РЗН 2014/2156"),
        ("сумка санитарная", None if "тактическ" in n else "ФСР 2009/05719"),
        ("сельских поселениях", "РЗН 2018/7096"),
        ("железнодорожном транспорте", "РЗН 2013/593"),
        ("аптечка фэст автомобильная", "ФСР 2010/06799"),
        ("первой помощи дорожная", "ФСР 2010/09275"),
        ("первой помощи транспортная", "ФСР 2010/09412"),
        ("для работников", "ФСР 2011/11668"),
        ("зс го", "ФСР 2009/05720") if "коллективная" in n else ("__", None),
        ("коллективная", "ФСР 2008/03785"),
        ("промышленных предприятий", "ФСР 2008/02948"),
        ("рабочих кабинетов", "ФСР 2008/03275"),
        ("универсальная", "ФСР 2008/02949"),
        ("мамы и малыша", "ФСР 2008/01761"),
        ("общего назначения аон", "ФСР 2009/05148"),
        ("нефтяника", "ФСР 2009/05718"),
        ("энергетика", "ФСР 2009/05717"),
        ("антиспид", "ФСР 2008/03331"),
        ("аптечка фэст первой помощи противоожоговая", "ФСР 2007/01032"),
        ("устройство-маска", "ФСР 2007/00679"),
        ("носилки фэст бескаркасные тактические", "РЗН 2024/22742"),
        ("носилки фэст медицинские мягкие", "РЗН 2020/12675"),
        ("покрывало", "РЗН 2025/25343"),
        ("дозированной компрессией", "РЗН 2025/24565"),
        ("жгут фэст михайлова одноразовый", "РЗН 2024/24182"),
        ("венозный", "РЗН 2021/14159"),
        ("маска медицинская", "РЗН 2022/16528"),
        ("бинт эластичный", "ФСР 2009/05993"),
    ]
    for key, ru in rules:
        if key in n:
            return ru
    return None


def _aptechka_type(section: Optional[str], name: str) -> Optional[str]:
    """«Тип аптечки» — назначение: по разделу прайса, с уточнением по названию."""
    low = (name or "").lower()
    if "антишок" in low or "посиндромн" in low:
        return "Бытовая"
    if any(k in low for k in ("перевозки опасных грузов", "для строителей", "удаленной промышленной")):
        return "Производственная"
    for prefix, kind in _TYPE_BY_SECTION:
        if (section or "").startswith(prefix):
            return kind
    return None


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _num(value) -> Optional[float]:
    try:
        return round(float(str(value).replace(",", ".").replace(" ", "")), 2)
    except (TypeError, ValueError):
        return None


def _fmt_int(value) -> Optional[str]:
    n = _num(value)
    if n is None:
        return None
    return str(int(n)) if n == int(n) else str(n).replace(".", ",")


def _split_name(cell: str):
    """Название и штрихкод из ячейки номенклатуры."""
    barcode = None
    name_lines = []
    for line in cell.replace("\r", "").split("\n"):
        line = line.strip()
        m = re.match(r"штрих\s*-?\s*код\s*:?\s*(\d+)", line, re.I)
        if m:
            barcode = m.group(1)
        elif line:
            name_lines.append(line)
    return re.sub(r"\s+", " ", " ".join(name_lines)).strip(), barcode


class FestParser(BaseParser):
    display_name = "ФЭСТ (аптечки, прайс .xlsx)"

    def parse(self, file_path: str) -> List[ProductIn]:
        wb = _open_workbook(file_path)
        ws = wb.worksheets[0]
        products: List[ProductIn] = []
        seen = set()
        section = None
        for row in ws.iter_rows(values_only=True):
            row = tuple(row) + (None,) * (COL_BOX + 1 - len(row))
            article, cell = _text(row[COL_ARTICLE]), _text(row[COL_NAME])
            if not article and not cell:
                head = _text(row[COL_SECTION])
                if head and not head.lower().startswith("прайс"):
                    section = re.sub(r"^\d+\.\s*", "", head).strip()
                continue
            if article.lower() == "артикул" or not article:
                continue  # повторяющаяся шапка таблицы
            if article in seen:
                continue
            seen.add(article)
            name, barcode = _split_name(cell)
            attrs: Dict[str, str] = {}
            if barcode:
                attrs["Штрихкод"] = barcode
            # «Ставка НДС» из прайса в таблицу не выводим (решение пользователя, 2026-10-06).
            pack = _fmt_int(row[COL_PACK])
            if pack:
                attrs["Количество в коробке"] = f"{pack} шт"
            box = _text(row[COL_BOX])
            if box:
                parts = [p.strip() for p in box.replace("\n", " ").split("/")]
                if len(parts) == 2 and parts[0] and parts[1]:
                    attrs["Объем коробки, м3"] = parts[0]
                    attrs["Вес коробки, кг"] = parts[1]
            cost = _num(row[COL_PRICE])
            extra = _EXTRA.get(article) or {}
            if extra.get("attributes"):
                attrs = {**extra["attributes"], **attrs}
            attrs.update(_REGISTRY.get(article, {}))
            # У донора «Тип» — это исполнение (переносной/стационарный); «Тип аптечки» — назначение.
            donor_type = attrs.pop("Тип аптечки", None)
            if donor_type and set(donor_type.split("##")) <= {"Переносной", "Портативный", "Стационарный"}:
                attrs["Исполнение"] = donor_type
            ru = _rzn_number(extra.get("name") or name, attrs.get("ТУ", ""))
            if ru:
                attrs["Реестр Росздравнадзор"] = ru
            if attrs.get("Товар") in ("Аптечка", "Укладка"):
                kind = _aptechka_type(section, extra.get("name") or name)
                if kind:
                    attrs["Тип аптечки"] = kind
            products.append(
                ProductIn(
                    article=article,
                    name=extra.get("name") or name or article,
                    brand=BRAND,
                    category=section,
                    cost_price=cost,
                    sale_price=float(int(cost * MARKUP)) if cost is not None else None,  # копейки отрезаются
                    image_url=" ".join(extra["images"]) if extra.get("images") else None,
                    external_variant_id=_VARIANT_IDS.get(article),
                    description=extra.get("description"),
                    attributes=attrs,
                )
            )
        wb.close()
        return products
