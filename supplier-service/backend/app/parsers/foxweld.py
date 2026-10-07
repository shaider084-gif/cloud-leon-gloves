"""Поставщик FoxWeld (slug foxweld, https://foxweld.ru).

Парсер понимает два формата файла:

1. «Итоговая таблица» (формат самого сервиса, как у Техноавиа): заголовки «Название товара или услуги»,
   «Артикул», «Изображения», «Цена продажи», «Описание» и набор «Параметр: X».
   Из неё берётся база товаров (1897 позиций: описания, фото, характеристики).
   Заголовки пользователя приводятся к стандартным столбцам сервиса:
   «Вес, кг» -> «Вес», «Объем, м3» -> «Размер, м3», «Реестр минпромторг» -> «Реестр Минпромторг».

2. Прайс-лист поставщика (три таблицы в Яндекс.Документах): строка заголовка с «Артикул», «Наименование»,
   «VIP, руб», «Розн., руб» (+ «Гарантия», «Краткое описание»), между ними строки-разделы
   («Электросварка > Аппараты …»). Читает parse_price_rows() — его же использует кнопка «Обновить»
   (foxweld_sync.py). Цены прайса пишутся в параметры «РРЦ» (розничная) и «Цена VIP»;
   закупка (cost_price) пока не заполняется — правило ещё не задано пользователем.
"""
import re
from typing import Dict, List, Optional

import openpyxl

from ..schemas import ProductIn
from .base import BaseParser

PARAM_PREFIX = "Параметр: "
BRAND_HEADER = "Параметр: Бренд"

# Заголовки итоговой таблицы пользователя -> стандартные ключи параметров сервиса.
PARAM_RENAME = {
    "Вес, кг": "Вес",
    "Объем, м3": "Размер, м3",
    "Реестр минпромторг": "Реестр Минпромторг",
}


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    text = str(value).strip().replace("\xa0", " ")
    return text or None


def _num(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _fmt_price(x: Optional[float]) -> Optional[str]:
    if x is None:
        return None
    return str(int(x)) if x == int(x) else ("%.2f" % x).rstrip("0").rstrip(".")


# ---------------------------------------------------------------- прайс-лист (источники «Обновить»)

def _find_columns(row) -> Optional[Dict[str, int]]:
    """Строка заголовка прайса: возвращает {поле: индекс столбца} или None."""
    cols: Dict[str, int] = {}
    for i, v in enumerate(row):
        t = str(v or "").strip().lower()
        if not t:
            continue
        if t.startswith("артикул") or t == "код" or t.startswith("код для заказа"):
            cols.setdefault("article", i)
        elif t.startswith("наименование") or t.startswith("название"):
            cols.setdefault("name", i)
        elif "vip" in t:
            cols.setdefault("vip", i)
        elif t.startswith("розн") or t.startswith("ррц") or "розничная" in t:
            cols.setdefault("retail", i)
        elif t.startswith("гаранти"):
            cols.setdefault("warranty", i)
        elif t.startswith("краткое описание") or t.startswith("описание"):
            cols.setdefault("brief", i)
    if "article" in cols and "name" in cols and ("retail" in cols or "vip" in cols):
        return cols
    return None


def parse_price_rows(file_path: str) -> List[Dict]:
    """Все листы файла-прайса -> [{article, name, vip, retail, warranty, brief, section}]."""
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    out: List[Dict] = []
    try:
        for ws in wb.worksheets:
            cols = None
            section = None
            for row in ws.iter_rows(values_only=True):
                if cols is None:
                    cols = _find_columns(row)
                    continue
                if _find_columns(row):  # повторная шапка таблицы внутри листа
                    cols = _find_columns(row)
                    continue

                def at(key):
                    i = cols.get(key)
                    return row[i] if i is not None and i < len(row) else None

                art = _clean(at("article"))
                name = _clean(at("name"))
                vip, retail = _num(at("vip")), _num(at("retail"))
                if not art:
                    first = _clean(row[0]) if row else None
                    if first and ">" in first and not name:  # «Электросварка > Аппараты …» — раздел
                        section = first
                    continue
                if not name or (vip is None and retail is None):
                    continue
                out.append({
                    "article": art,
                    "name": name,
                    "vip": vip,
                    "retail": retail,
                    "warranty": _clean(at("warranty")),
                    "brief": _clean(at("brief")),
                    "section": section,
                })
    finally:
        wb.close()
    return out


def price_attributes(row: Dict) -> Dict[str, str]:
    """Параметры из строки прайса, которые обновляются при каждой синхронизации."""
    attrs: Dict[str, str] = {}
    if row.get("retail") is not None:
        attrs["РРЦ"] = _fmt_price(row["retail"])
    if row.get("vip") is not None:
        attrs["Цена VIP"] = _fmt_price(row["vip"])
    return attrs


# ---------------------------------------------------------------- итоговая таблица (база)

class FoxweldParser(BaseParser):
    display_name = "FoxWeld (итоговая таблица или прайс-лист)"

    def parse(self, file_path: str) -> List[ProductIn]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]

        if "Название товара или услуги" not in headers:
            # формат прайс-листа поставщика (Яндекс-таблица)
            return [
                ProductIn(
                    article=r["article"], name=r["name"], description=r.get("brief"),
                    attributes={**price_attributes(r), **({"Гарантия, мес": r["warranty"]} if r.get("warranty") else {})},
                )
                for r in parse_price_rows(file_path)
            ]

        products: List[ProductIn] = []
        for row in rows[1:]:
            if row is None or all(v is None for v in row):
                continue
            values = dict(zip(headers, row))
            name = _clean(values.get("Название товара или услуги"))
            if not name:
                continue
            attrs: Dict[str, str] = {}
            brand = None
            for header, value in values.items():
                cleaned = _clean(value)
                if not cleaned:
                    continue
                if header == BRAND_HEADER:
                    brand = cleaned
                elif header.startswith(PARAM_PREFIX):
                    key = header[len(PARAM_PREFIX):]
                    attrs[PARAM_RENAME.get(key, key)] = cleaned
                elif header in PARAM_RENAME:  # «Вес, кг» без префикса
                    attrs[PARAM_RENAME[header]] = cleaned
            if attrs.get("Вес"):  # как везде в проекте — десятичная запятая
                attrs["Вес"] = attrs["Вес"].replace(".", ",")
            products.append(
                ProductIn(
                    article=_clean(values.get("Артикул")),
                    name=name,
                    brand=brand,
                    sale_price=_num(values.get("Цена продажи")),
                    image_url=_clean(values.get("Изображения")),
                    description=_clean(values.get("Описание")),
                    attributes=attrs,
                )
            )
        return products
