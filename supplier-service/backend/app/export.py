from io import BytesIO
from typing import List

import openpyxl
from openpyxl.styles import Font, PatternFill

from . import models

FIXED_COLUMNS = [
    ("article", "Артикул"),
    ("name", "Название"),
    ("brand", "Бренд"),
    ("category", "Категория"),
    ("unit", "Ед. изм."),
    ("cost_price", "Закупочная цена"),
    ("sale_price", "Цена продажи"),
]

HEADER_FILL = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")


def export_products_to_xlsx(products: List[models.Product]) -> BytesIO:
    """Собирает products в xlsx: фиксированные колонки + динамические 'Параметр: X'
    из объединения ключей attributes по всем товарам (аналог структуры
    'Шаблон сиз рук')."""

    attribute_keys: List[str] = []
    seen = set()
    for p in products:
        for key in (p.attributes or {}).keys():
            if key not in seen:
                seen.add(key)
                attribute_keys.append(key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары"

    headers = [label for _, label in FIXED_COLUMNS] + [f"Параметр: {k}" for k in attribute_keys]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for p in products:
        row = [getattr(p, field) for field, _ in FIXED_COLUMNS]
        row += [(p.attributes or {}).get(k, "") for k in attribute_keys]
        ws.append(row)

    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 60)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
